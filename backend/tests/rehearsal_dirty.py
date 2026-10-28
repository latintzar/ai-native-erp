"""rehearsal_dirty.py — M1M2 WP-2.5: THE EARLY DIRTY SMOKE REHEARSAL (CEO finding 4).

Not a vital — a REPEATABLE rehearsal that drives a genuinely UGLY, real-world-shaped price list
through the REAL pipeline on the rehearsal clone, hunting structural unknowns after 2 WPs instead
of 6 (WP-7 stays the full seal):

  A. FIXTURES — an ugly CSV tarifa (accents, merged-cell vibes, section rows, duplicate headers,
     «docena/sarta/manojo/bandeja» units, comma-decimal € prices, missing prices, missing codes,
     footers) + the same sheet as a hand-built PDF, + a photo (PIL-rendered JPEG) when PIL exists.
  B. PARSE — the real wizard path: onboarding_ext.process_catalogue → skuparse (REAL Sonnet
     extraction + Haiku legit-check) → wizard_state.skus. PDF and photo parsed via
     skuparse.parse_catalogue directly (the same dispatch skuparse runs in production).
  C. GOLIVE — the real golive.activate: WP-1 seed (embed+upsert), tenant live + number bound,
     lifecycle derives onboarded, welcome captured (egress-sandboxed + belt).
  D. ORDER — a dirty es-ES chat order through the REAL 5-scale pipeline
     (brain.pipeline.run_message persist=True) grounding against the seeded catalogue.
  E. NUDGE — the order left unreviewed at the cutoff: _review_nudge_scan emits
     orders.review_pending ONCE; the armed r1_nudge rung FIRES a WhatsApp (gate-checked);
     a scheduled dispatch with nothing approved sends the held-only notice.
  F. DISPATCH — approve (submit_order) → _run_dispatch: the REAL Excel bytes carry THEIR columns,
     THEIR catalogue codes (codigo) and the customer's codigo_cliente; the dirty customer name
     gets a sanitized sheet title; the value snapshot freezes; the review tab RESOLVES.

Every rough edge observed lands in the ROUGH list, printed at the end — the whole point of WP-2.5.

SAFETY: hard-refuses the live `rocoto` DB and refuses to run without ROCOTO_EGRESS_SANDBOX=1.
LLM calls (Anthropic) are REAL — that is the rehearsal's purpose. The heavy Opus `understand`
pass is stubbed by default (not on the smoke chain); set REHEARSAL_UNDERSTAND=1 to run it.
REHEARSAL_KEEP=1 leaves all rows + artifacts in place for inspection.

Run on the clone (prod box):
  source deploy/harness/rehearsal_env.sh
  PYTHONPATH=. /opt/rocoto/venv/bin/python tests/rehearsal_dirty.py
"""
import asyncio
import io
import os
import re
import sys
import tempfile
from datetime import datetime, timezone
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

_DSN = os.environ.get("REHEARSAL_DSN")
if _DSN:
    for _k in ("DATABASE_URL", "RUNTIME_DATABASE_URL", "CONTROL_DATABASE_URL"):
        os.environ.setdefault(_k, _DSN)

for _l in Path("/opt/rocoto/.env").read_text().splitlines() if Path("/opt/rocoto/.env").exists() else []:
    _l = _l.strip()
    if _l and not _l.startswith("#") and "=" in _l:
        _k, _, _v = _l.partition("=")
        os.environ.setdefault(_k.strip(), _v.strip())

# the class-E floor is a PRECONDITION here, not a default — refuse to run without it.
if os.environ.get("ROCOTO_EGRESS_SANDBOX") not in ("1", "true", "True"):
    raise SystemExit("REFUSING: rehearsal_dirty needs ROCOTO_EGRESS_SANDBOX=1 (source rehearsal_env.sh)")

import app as app_mod  # noqa: E402
import db  # noqa: E402
import directory  # noqa: E402
import duties  # noqa: E402
import orders as orders_mod  # noqa: E402
import presence  # noqa: E402
import reach  # noqa: E402
import runtime  # noqa: E402
import spine  # noqa: E402
import twilio_wa  # noqa: E402
import worker  # noqa: E402
from brain import events as brain_events  # noqa: E402
from brain import (  # noqa: E402
    golive,
    matcher,
    skuparse,  # noqa: E402
)
from brain import onboarding_ext as ob_ext  # noqa: E402
from brain import pipeline as brain_pipeline  # noqa: E402

LEAD_ID = "rehdirty_lead"
TOKEN = "rehdirty_tok"
BOSS_PHONE = "+34600717288"                    # the supplier (boss) — rehearsal-reserved
BOSS_SCOPE = runtime.scope_of(BOSS_PHONE)
BOSS_NAME = "Amparo Zurdo"
BUSINESS = "Cárnicas El Zurdo e Hijos, S.L."   # accents + punctuation → the derived-slug path
CUST_PHONE = "+34600717299"                    # the dirty customer
CUST_SCOPE = runtime.scope_of(CUST_PHONE)
CUST_NAME = 'Bar "El Niño" / Q:*'              # openpyxl-illegal sheet-title chars on purpose
BIND_NUMBER = "+34698765432"
ERP_CODE = "CL-0077"
KEEP = os.environ.get("REHEARSAL_KEEP") in ("1", "true")
_ILLEGAL = set('\\/*?:[]')

R, ROUGH = [], []


def ck(name, ok, detail=""):
    R.append(bool(ok))
    print(f"  [{'PASS' if ok else 'FAIL'}] {name}" + (f" — {detail}" if detail else ""))


def rough(note):
    ROUGH.append(note)
    print(f"  [ROUGH] {note}")


# ═══ A · the ugly fixtures ═══════════════════════════════════════════════════════════════════════

DIRTY_ROWS = [
    "TARIFA PRECIOS 2026 — CÁRNICAS EL ZURDO E HIJOS S.L.;;;;",
    "C/ Matadero Viejo 12, Cuenca · Tel 969 22 11 00;;;;",
    ";;;;",
    "REF;PRODUCTO;FORMATO;PRECIO;OBS",
    "   EMBUTIDOS   ;;;;",
    "CH-01;Chorizo picante casero (sarta);sarta;3,80 €;",
    "CH-02;CHORIZO DULCE  EXTRA;sarta;3,60;",
    "MO-11;Morcilla de cebolla;docena;14,50 €;pedido mínimo 2 doc",
    "SA-3; Salchichón ibérico VELA ;kg;11,90 €;",
    ";Jamón curado bodega (pieza entera);pieza;;precio según peso",
    "QUESOS;;;;",
    "QM/12;Queso manchego semicurado;ud;8,75 €;cuña 250 g aparte",
    "QM/13;Queso manchego CURADO reserva;ud;12,40 €;",
    ";Queso de cabra rulo;ud;6,20 €;",
    "CONSERVAS Y VARIOS;;;;",
    "TO-77;Tomate seco en aceite (tarro 314 ml);caja;28,80 €;12 uds/caja",
    ";Huevos camperos XL;docena;2,95 €;",
    "AC-9;Aceite oliva virgen extra 5 L;caja;;consultar",
    "MJ-2;Manojo ajos tiernos;manojo;1,80 €;temporada",
    "REF;PRODUCTO;FORMATO;PRECIO;OBS",
    "BA-1;Bandeja pinchos morunos (12 uds);bandeja;7,20 €;",
    "* Precios sin IVA;;;;",
    "IVA no incluido — portes pagados a partir de 150 €;;;;",
]
N_REAL_PRODUCTS = 13     # what a careful human counts in the sheet above


def _write_csv(d: str) -> str:
    p = os.path.join(d, "tarifas FINAL2 (1).csv")
    Path(p).write_text("\n".join(DIRTY_ROWS), encoding="utf-8")
    return p


def _pdf_escape(s: str) -> bytes:
    # WinAnsi (latin-1) text with () and \ escaped — enough for a hand-built price sheet.
    b = s.encode("latin-1", "replace")
    return b.replace(b"\\", b"\\\\").replace(b"(", b"\\(").replace(b")", b"\\)")


def _write_pdf(d: str) -> str:
    """A REAL (minimal, valid) one-page PDF of the same tarifa — exercises the pypdf branch of
    skuparse for real, no fixture monkeypatching."""
    lines = [r.replace(";", "   ") for r in DIRTY_ROWS]
    content = b"BT /F1 10 Tf 40 800 Td 13 TL\n"
    for ln in lines:
        content += b"(" + _pdf_escape(ln) + b") Tj T*\n"
    content += b"ET\n"
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
         b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"endstream",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, o in enumerate(objs, start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + o + b"\nendobj\n")
    xref_at = out.tell()
    out.write(b"xref\n0 " + str(len(objs) + 1).encode() + b"\n0000000000 65535 f \n")
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(b"trailer\n<< /Size " + str(len(objs) + 1).encode() +
              b" /Root 1 0 R >>\nstartxref\n" + str(xref_at).encode() + b"\n%%EOF\n")
    p = os.path.join(d, "tarifa_escaneada.pdf")
    Path(p).write_bytes(out.getvalue())
    return p


def _write_photo(d: str):
    """A JPEG 'phone photo' of the tarifa (PIL-rendered, slight rotation + off-white paper).
    None when PIL is missing — the photo path is then reported as SKIPPED, not failed."""
    try:
        from PIL import Image, ImageDraw
    except Exception:
        return None
    img = Image.new("RGB", (1240, 1000), (246, 242, 233))
    dr = ImageDraw.Draw(img)
    y = 30
    for ln in DIRTY_ROWS:
        dr.text((40, y), ln.replace(";", "   "), fill=(28, 24, 20))
        y += 40
    img = img.rotate(1.6, expand=True, fillcolor=(210, 205, 196))
    p = os.path.join(d, "IMG_4123.jpg")
    img.save(p, "JPEG", quality=68)
    return p


# ═══ helpers ══════════════════════════════════════════════════════════════════════════════════════

def _wb(xlsx: bytes):
    from openpyxl import load_workbook
    return load_workbook(io.BytesIO(xlsx))


def _headers(ws):
    out = []
    for i in range(1, 25):
        v = ws.cell(row=1, column=i).value
        if v is None:
            break
        out.append(v)
    return out


async def _events_on(scope, etype):
    async with db.control_pool().acquire() as con:
        return int(await con.fetchval(
            "select count(*) from event where scope=$1::ltree and type=$2", scope, etype) or 0)


async def _cleanup(tenant=None):
    async with db.get_pool().acquire() as con:
        if tenant:
            await con.execute("delete from sku_aliases where tenant_id=$1", tenant)
            await con.execute("delete from catalogue_skus where tenant_id=$1", tenant)
            await con.execute("delete from catalogue_directory where tenant_id=$1", tenant)
            await con.execute("delete from tenants where slug=$1", tenant)
        await con.execute("delete from onboarding_progress where token=$1", TOKEN)
        await con.execute("delete from onboarding_tokens where token=$1", TOKEN)
        await con.execute("delete from leads where id=$1", LEAD_ID)
        await con.execute("delete from orders where phone_number = any($1::text[])",
                          [CUST_PHONE, BOSS_PHONE])
        await con.execute("delete from customer_profiles where phone_number = any($1::text[])",
                          [CUST_PHONE, BOSS_PHONE])
        await con.execute("delete from whatsapp_messages where phone_number = any($1::text[])",
                          [CUST_PHONE, BOSS_PHONE])
        await con.execute("delete from agent_settings where agent_id='orders'")
        await con.execute("update manifest_reflex set config = config - 'boss_phone' "
                          "where organ = any($1::text[]) and scope_tier='boss'",
                          [["review_nudge", "dispatch_failed"]])
    if tenant:
        task = directory._PENDING.pop(tenant, None)
        if task and not task.done():
            task.cancel()
        try:
            async with db.scoped_tx(f"t.{tenant}") as con:
                await con.execute("delete from event where scope=$1::ltree", f"t.{tenant}")
        except Exception:
            pass
    for sc in (BOSS_SCOPE, CUST_SCOPE):
        try:
            async with db.scoped_tx(sc) as con:
                for t in ("timer", "open_loop", "action", "event"):
                    await con.execute(f"delete from {t} where scope=$1::ltree", sc)
                await con.execute("delete from scope_node where scope=$1::ltree", sc)
        except Exception:
            pass


async def main():
    await db.init()
    duties.invalidate_cache()

    async with db.get_pool().acquire() as con:
        dbname = await con.fetchval("select current_database()")
    if dbname == "rocoto":
        raise SystemExit("REFUSING to rehearse against the LIVE 'rocoto' DB — use rocoto_rehearsal")
    print(f"  (db={dbname}, egress=SANDBOX)")
    async with db.get_pool().acquire() as con:
        await directory.ensure_schema(con)

    await _cleanup()

    tmpdir = tempfile.mkdtemp(prefix="rehdirty_")
    csv_path = _write_csv(tmpdir)
    pdf_path = _write_pdf(tmpdir)
    photo_path = _write_photo(tmpdir)
    print(f"  (fixtures in {tmpdir})")

    # the wizard-shaped lead: paid, mid-wizard (token ACTIVE — that's when catalogues upload)
    async with db.get_pool().acquire() as con:
        await con.execute("insert into leads (id, phone, name, business, email) values ($1,$2,$3,$4,$5)",
                          LEAD_ID, BOSS_PHONE, BOSS_NAME, BUSINESS, "rehearsal@rehearsal.local")
        await con.execute("insert into onboarding_tokens (token, lead_id, status) values ($1,$2,'active')",
                          TOKEN, LEAD_ID)
        await con.execute(
            "insert into onboarding_progress (token, wizard_state, step, version, paid_at) "
            "values ($1,$2,3,1,$3)", TOKEN, {}, datetime.now(timezone.utc))

    # ═══ B · PARSE — the real wizard catalogue path on the ugly CSV ═══════════════════════════════
    _understood = []

    async def _stub_understand(token):
        _understood.append(token)
    _real_ubg = ob_ext._understand_bg
    if os.environ.get("REHEARSAL_UNDERSTAND") not in ("1", "true"):
        ob_ext._understand_bg = _stub_understand   # the heavy Opus pass is off the smoke chain

    t0 = asyncio.get_event_loop().time()
    parsed = await ob_ext.process_catalogue(TOKEN, csv_path, os.path.basename(csv_path))
    dt_csv = asyncio.get_event_loop().time() - t0
    ck("CSV: process_catalogue ok + legit on the ugly tarifa",
       parsed.get("ok") and parsed.get("legit"), f"{parsed} ({dt_csv:.1f}s)")
    async with db.get_pool().acquire() as con:
        state = await con.fetchval("select wizard_state from onboarding_progress where token=$1", TOKEN)
    import json as _json
    state = state if isinstance(state, dict) else _json.loads(state or "{}")
    skus = state.get("skus") or []
    n = len(skus)
    by_name = { (s.get("name") or "").lower(): s for s in skus }

    def find(frag):
        return next((s for k, s in by_name.items() if frag in k), None)

    ck(f"CSV: plausible product count (human count {N_REAL_PRODUCTS})", 10 <= n <= 16, f"n={n}")
    if n != N_REAL_PRODUCTS:
        rough(f"parser count {n} vs human count {N_REAL_PRODUCTS} — diff rows: check section/footer bleed")
    junk = [s for s in skus if re.search(r"\bIVA\b|TARIFA|\bTel\b|^\s*REF\s*$|^EMBUTIDOS|^QUESOS|^CONSERVAS",
                                         (s.get("name") or ""), re.I)]
    ck("CSV: no section headers / footers minted as products", not junk,
       f"junk={[s.get('name') for s in junk]}")
    cho = find("chorizo picante")
    ck("CSV: code survives (CH-01 on the chorizo)", bool(cho) and (cho.get("code") == "CH-01"), str(cho))
    qm = find("semicurado")
    ck("CSV: slashed code QM/12 survives verbatim", bool(qm) and qm.get("code") == "QM/12", str(qm))
    ck("CSV: comma-decimal € price parsed (chorizo 3.80)",
       bool(cho) and abs(float(cho.get("price") or 0) - 3.8) < 0.005, str(cho and cho.get("price")))
    jam = find("jamón") or find("jamon")
    ck("CSV: missing price stays null — never invented (jamón)",
       jam is not None and jam.get("price") in (None, ""), str(jam))
    hue = find("huevos")
    ck("CSV: codeless row keeps empty code + docena unit (huevos)",
       hue is not None and not (hue.get("code") or "") and "docena" in (hue.get("unit") or "").lower(),
       str(hue))
    weird_units = {s.get("unit") for s in skus if (s.get("unit") or "").lower()
                   not in ("kg", "ud", "caja", "unidad", "")}
    if weird_units:
        print(f"  [info] non-canonical units preserved as written: {sorted(weird_units)}")

    # PDF — the same tarifa through the real pypdf branch
    t0 = asyncio.get_event_loop().time()
    pdfres = await skuparse.parse_catalogue(pdf_path, "tarifa_escaneada.pdf")
    dt_pdf = asyncio.get_event_loop().time() - t0
    npdf = pdfres.get("count", 0)
    ck("PDF: real pypdf branch extracts the tarifa (count within ±3 of CSV)",
       pdfres.get("ok") and abs(npdf - n) <= 3, f"count={npdf} vs csv={n} ({dt_pdf:.1f}s)")
    if pdfres.get("ok") and npdf != n:
        rough(f"PDF parse count {npdf} != CSV count {n} — same sheet, different reading")

    # PHOTO — real Claude-vision OCR when PIL could render one
    if photo_path:
        t0 = asyncio.get_event_loop().time()
        phres = await skuparse.parse_catalogue(photo_path, "IMG_4123.jpg")
        dt_ph = asyncio.get_event_loop().time() - t0
        nph = phres.get("count", 0)
        ck("PHOTO: OCR branch returns products from the rendered snapshot",
           phres.get("ok") and nph >= 8, f"count={nph} ({dt_ph:.1f}s)")
        if phres.get("ok") and abs(nph - n) > 2:
            rough(f"photo OCR count {nph} vs CSV {n} — vision path loses/gains rows on a clean render; "
                  f"a real crumpled photo will be worse")
    else:
        print("  [SKIP] photo path (PIL not installed in this venv)")
        rough("photo fixture SKIPPED — PIL missing; the OCR branch went unexercised in this run")

    ob_ext._understand_bg = _real_ubg
    if _understood:
        print(f"  [info] heavy Opus understand pass stubbed for token(s) {_understood} "
              f"(REHEARSAL_UNDERSTAND=1 to run it)")

    # ═══ C · GOLIVE — the real activate on the parsed mess ════════════════════════════════════════
    SENT = []
    _real_send = twilio_wa.send_whatsapp

    async def _capture_send(phone, text, media_url=None):
        SENT.append((phone, text))
        return "REHEARSAL-wa"
    twilio_wa.send_whatsapp = _capture_send

    g = await golive.activate(BOSS_SCOPE, BIND_NUMBER)
    tenant = g.get("tenant") or ""
    ck("GOLIVE: activate ok, lifecycle derives onboarded",
       g.get("ok") and not g.get("already_live") and g.get("lifecycle") == "onboarded",
       str({k: g.get(k) for k in ('ok', 'reason', 'lifecycle', 'tenant')}))
    ck("GOLIVE: derived tenant slug is ltree-safe (accents/S.L. folded)",
       bool(re.fullmatch(r"[a-z0-9_]+", tenant or "-")), f"tenant={tenant}")
    async with db.get_pool().acquire() as con:
        trow = await con.fetchrow("select status, number, golive_at from tenants where slug=$1", tenant)
        n_cat = int(await con.fetchval(
            "select count(*) from catalogue_skus where tenant_id=$1 and in_assortment", tenant) or 0)
        n_emb = int(await con.fetchval(
            "select count(*) from catalogue_skus where tenant_id=$1 and embedding is not null", tenant) or 0)
    ck("GOLIVE: tenant live + number bound + catalogue fully embedded",
       bool(trow) and trow["status"] == "live" and trow["number"] == BIND_NUMBER
       and n_cat == n and n_emb == n_cat,
       f"cat={n_cat}/{n} embedded={n_emb}")
    ck("GOLIVE: welcome sent once to the boss (egress-sandboxed + captured)",
       len(SENT) == 1 and SENT[0][0] == BOSS_PHONE and "Rocoto" in SENT[0][1], f"sends={len(SENT)}")

    # Run #1 found the window (an order arriving inside the 20s debounce grounds with zero
    # candidates); golive now FORCES the distill synchronously on activation (brain/golive.py (3b)).
    # So this call is the steady-state RE-CHECK: the directory must be present and complete whether
    # or not THIS call rebuilt it — `unchanged` means activation already did, which is the fix
    # working, not a miss. Asserting only `built` here asserted the fix had NOT landed.
    db_dir = await directory.build(tenant)
    ck("DIRECTORY: distill present + complete from the seeded mess",
       (bool(db_dir.get("built")) or db_dir.get("reason") == "unchanged")
       and int(db_dir.get("n_skus") or 0) == n, str(db_dir))
    rough("the ACTIVATION path now forces the distill synchronously, but a later catalogue edit still "
          "rides the ~20s debounce (lost on a process restart) — an order inside that window sees a "
          "stale directory")

    m = await matcher.match("chorizo picante", tenant=tenant)
    ck("GROUND: full dirty mention grounds to CH-01 (the pipeline's shape)",
       ((await matcher.match("sartas de chorizo picante", tenant=tenant)).pick or {}).get("code") == "CH-01")
    if (m.pick or {}).get("code") != "CH-01":
        rough(f"short mention «chorizo picante» reranks to {(m.pick or {}).get('name')!r} "
              f"(band={m.band}) — the bge cross-encoder misses the picante/dulce distinction on "
              f"short queries; full mentions pick right")
    m2 = await matcher.match("queso curao", tenant=tenant)
    print(f"  [info] «queso curao» → band={m2.band} pick={(m2.pick or {}).get('name')!r} "
          f"(semicurado/curado ambiguity — expected sub-clean)")
    if (m2.pick or {}).get("name") and "curado" not in (m2.pick or {}).get("name", "").lower():
        rough(f"«queso curao» grounded to {(m2.pick or {}).get('name')!r} — ambiguity resolution off")

    # ═══ D · the dirty chat order through the REAL pipeline ═══════════════════════════════════════
    async with db.scoped_tx(CUST_SCOPE) as con:
        await spine.resolve_chat_scope(con, None, "wa", CUST_PHONE, display_name=CUST_NAME,
                                       scope_override=CUST_SCOPE)
    # the brain's identity is customer_profiles, NOT the spine contact (run #1 finding: with no
    # profile, the synthesiser minted the customer name from the greeting text — «Amparo (CL-0077)»).
    async with db.get_pool().acquire() as con:
        await con.execute(
            "insert into customer_profiles (phone_number, tenant_id, display_name) values ($1,$2,$3) "
            "on conflict (tenant_id, phone_number) do update set tenant_id=$2, display_name=$3",
            CUST_PHONE, tenant, CUST_NAME)
    rough("identity is TWO systems: spine contact display_name does NOT feed brain identity "
          "(customer_profiles) — an unprofiled customer got a name invented from the greeting "
          "text («Amparo (CL-0077)») in run #1")
    ORDER_TEXT = ("buenas amparo!! pa mañana ponme 2 sartas de chorizo picante, media docena huevos... "
                  "no q va, UNA docena mejor. 3 cajas del tomate seco y 2 del queso manchego curado. "
                  "ah y apunta mi numero de cliente q es CL-0077, siempre lo pones mal jeje")
    t0 = asyncio.get_event_loop().time()
    order = await brain_pipeline.run_message(ORDER_TEXT, CUST_PHONE, persist=True, tenant=tenant)
    dt_ord = asyncio.get_event_loop().time() - t0
    ck("ORDER: the 5-scale pipeline builds an order from the dirty burst",
       bool(order) and bool(order.get("lines")),
       f"{len((order or {}).get('lines') or [])} lines ({dt_ord:.1f}s)")
    oid = (order or {}).get("id")
    lines = (order or {}).get("lines") or []
    for ln in lines:
        print(f"    line: {ln.get('item')!r} qty={ln.get('qty')} {ln.get('unit')} "
              f"code={ln.get('code')!r} price={ln.get('price')} band={ln.get('band')} "
              f"sku_id={ln.get('sku_id')!r}")
    ck("ORDER: persisted needs_review (the review gate holds)",
       (order or {}).get("status", "needs_review") == "needs_review", str((order or {}).get("status")))
    txt = _json.dumps(lines, ensure_ascii=False).lower()
    ck("ORDER: the self-correction stuck (12 huevos, not 6)",
       ("huevo" in txt) and not re.search(r'"qty":\s*0?\.5\b', txt),
       "docena self-correct")
    got_codes = {ln.get("code") for ln in lines if ln.get("code")}
    ck("ORDER: grounded lines carry THEIR catalogue codes", bool(got_codes & {"CH-01", "TO-77", "QM/13"}),
       f"codes={sorted(got_codes)}")
    if not any((ln.get("code") or "") == "QM/13" for ln in lines):
        rough("the «queso manchego curado» line did not ground to QM/13 — check which SKU it picked")

    # the customer's ERP client code — the WP-3 conversational-capture write (db layer directly here)
    await db.erp_code_set(CUST_PHONE, ERP_CODE)

    # ═══ E · left unreviewed at the cutoff — the nudge chain, then the held-only email ════════════
    async with db.scoped_tx(BOSS_SCOPE) as con:
        await spine.resolve_chat_scope(con, None, "wa", BOSS_PHONE, display_name=BOSS_NAME,
                                       scope_override=BOSS_SCOPE, tier_override="boss")
        await con.execute("update scope_node set tier='boss' where scope=$1::ltree", BOSS_SCOPE)
    async with db.get_pool().acquire() as con:
        await con.execute("update manifest_reflex set config = config || $1::jsonb "
                          "where organ='review_nudge' and scope_tier='boss'", {"boss_phone": BOSS_PHONE})

    THEIR_COLS = [
        {"header": "REF", "field": "codigo", "order": 0},
        {"header": "Nº CLIENTE", "field": "codigo_cliente", "order": 1},
        {"header": "CLIENTE", "field": "cliente", "order": 2},
        {"header": "ARTÍCULO", "field": "producto", "order": 3},
        {"header": "CANT", "field": "cantidad", "order": 4},
        {"header": "FORMATO", "field": "unidad", "order": 5},
        {"header": "PVP", "field": "precio", "order": 6},
        {"header": "NOTAS", "field": "nota", "order": 7},
    ]
    CUST_COLS = [
        {"header": "Ref.", "field": "codigo", "order": 0},
        {"header": "Cód. cliente", "field": "codigo_cliente", "order": 1},
        {"header": "Artículo", "field": "producto", "order": 2},
        {"header": "Cant.", "field": "cantidad", "order": 3},
        {"header": "Formato", "field": "unidad", "order": 4},
    ]
    from datetime import timedelta
    now_real = datetime.now(presence.TZ).replace(second=0, microsecond=0)
    cutoff = now_real + timedelta(minutes=20)      # inside the 45-min lead window; a +30min pacing
    cutoff_hhmm = cutoff.strftime("%H:%M")         # hold would OVERSHOOT it → the exemption sends NOW
    if cutoff.date() != now_real.date():
        rough("cutoff crossed midnight — nudge sub-chain will no-op this run (rerun earlier in the day)")
    ST = {"mode": "scheduled", "schedule_time": cutoff_hhmm, "recipient_email": "rehearsal@rehearsal.local",
          "excel_format": "standard", "n": 10, "default_columns": THEIR_COLS,
          "per_customer_columns": {CUST_PHONE: CUST_COLS}}
    await db.settings_set("orders", ST)

    r1 = await app_mod._review_nudge_scan(ST, now=now_real)
    r2 = await app_mod._review_nudge_scan(ST, now=now_real)
    n_pend = await _events_on(BOSS_SCOPE, "orders.review_pending")
    ck("NUDGE: unreviewed at cutoff → orders.review_pending fires ONCE",
       r1.get("emitted") and not r2.get("emitted") and n_pend == 1,
       f"r1={r1} r2={r2} events={n_pend}")

    ekey = f"revnudge:{cutoff.strftime('%Y%m%d-%H%M')}"
    async with db.control_pool().acquire() as con:
        tab = await con.fetchrow("select id, contingencies from open_loop where scope=$1::ltree "
                                 "and entity_key=$2 and status='open'", BOSS_SCOPE, ekey)
        tid = await con.fetchval(
            "select t.id from timer t join open_loop l on l.id=t.loop_id where l.scope=$1::ltree "
            "and l.entity_key=$2 and t.branch_id='r1_nudge' and t.state='armed'", BOSS_SCOPE, ekey)
    cont = tab["contingencies"] if tab else "[]"
    cont = cont if isinstance(cont, list) else _json.loads(cont or "[]")
    sent0 = ((cont[0] if cont else {}).get("action", {}) or {}).get("sentence") or ""
    ck("NUDGE: tab open + r1_nudge armed + resolved count-free sentence",
       tab is not None and tid is not None and sent0 and "{" not in sent0 and cutoff_hhmm in sent0,
       sent0[:80])

    # fire the rung for real (worker path; gate.check_outbound runs) — verify/reach stubbed like the vital
    _nudge_sends = []

    async def _cap2(phone, text, media_url=None):
        _nudge_sends.append((phone, text))
        return "REHEARSAL-wa2"

    async def _fake_verify(con, **kw):
        return True

    async def _fake_reach(con, scope, intent="nudge", has_email=False):
        return {"channel": "freeform"}
    _real_verify, _real_reach = worker._verify_branch, reach.choose_channel
    worker._verify_branch, reach.choose_channel = _fake_verify, _fake_reach
    twilio_wa.send_whatsapp = _cap2
    fired = False
    if tid is not None:
        async with db.scoped_tx(BOSS_SCOPE) as con:
            fired = await worker._fire_one(con, {"id": int(tid), "scope": BOSS_SCOPE})
    worker._verify_branch, reach.choose_channel = _real_verify, _real_reach
    twilio_wa.send_whatsapp = _capture_send
    ck("NUDGE: r1_nudge FIRES the WhatsApp to the boss (spine outbound, gate-checked)",
       fired and len(_nudge_sends) == 1 and _nudge_sends[0][0] == BOSS_PHONE,
       (_nudge_sends[0][1][:80] if _nudge_sends else "no send"))

    # the daily cutoff with NOTHING approved → the held-only notice email (through the REAL rail,
    # which must sandbox-short-circuit — the WP-2.5 class-E fix under test in its real chain)
    _emails = []
    _real_email = orders_mod.send_email_async

    async def _spy_email(to, subject, html, attachment_name=None, attachment_bytes=None):
        res = await _real_email(to, subject, html,
                                attachment_name=attachment_name, attachment_bytes=attachment_bytes)
        _emails.append({"to": to, "subject": subject, "html": html,
                        "attachment": attachment_bytes, "res": res})
        return res
    orders_mod.send_email_async = _spy_email

    d0 = await app_mod._run_dispatch("scheduled")
    ck("CUTOFF: nothing approved + 1 held → the held-only notice goes out (never a silent day)",
       d0.get("ok") and d0.get("note") == "held-only notice" and len(_emails) == 1
       and "sin revisar" in _emails[0]["subject"],
       f"res={ {k: d0.get(k) for k in ('ok', 'note', 'held')} } "
       f"subject={_emails[0]['subject'] if _emails else None}")
    ck("CUTOFF: the real Resend rail short-circuited under the sandbox (class-E fix live)",
       bool(_emails) and _emails[0]["res"].get("id") == "SANDBOX-email", str(_emails and _emails[0]["res"]))

    # ═══ F · approve → dispatch → THEIR Excel ═════════════════════════════════════════════════════
    approved = await brain_events.submit_order(oid, actor="operator") if oid else None
    ck("APPROVE: submit_order flips the order to approved", (approved or {}).get("status") == "approved",
       str((approved or {}).get("status")))

    d1 = await app_mod._run_dispatch("scheduled")
    att = next((e["attachment"] for e in reversed(_emails) if e["attachment"]), None)
    ck("DISPATCH: ok, 1 order, 0 held, Excel attached",
       d1.get("ok") and d1.get("count") == 1 and d1.get("held") == 0 and att is not None,
       str({k: d1.get(k) for k in ('ok', 'count', 'held')}))

    if att:
        out_xlsx = os.path.join(tmpdir, "dispatched.xlsx")
        Path(out_xlsx).write_bytes(att)
        print(f"  [info] dispatched Excel saved → {out_xlsx}")
        wb = _wb(att)
        names = wb.sheetnames
        legal = all(not (_ILLEGAL & set(t)) and 0 < len(t) <= 31 for t in names)
        ck("EXCEL: dirty customer name → sanitized per-customer sheet title",
           legal and any("Niño" in t for t in names), f"sheets={names}")
        ws = wb[next((t for t in names if "Niño" in t), names[0])]
        hdrs = _headers(ws)
        ck("EXCEL: THEIR per-customer columns render exactly",
           hdrs == [c["header"] for c in CUST_COLS], f"headers={hdrs}")
        rows = []
        for ri in range(2, 2 + len(lines)):
            rows.append({h: ws.cell(row=ri, column=i + 1).value for i, h in enumerate(hdrs)})
        for r_ in rows:
            print(f"    excel row: {r_}")
        codes_in_xl = {r_.get("Ref.") for r_ in rows if r_.get("Ref.")}
        ck("EXCEL: codigo column carries THEIR catalogue codes",
           bool(codes_in_xl & {"CH-01", "TO-77", "QM/13", "QM/12"}), f"codes={sorted(codes_in_xl)}")
        ck("EXCEL: codigo_cliente resolves at dispatch (CL-0077 on every line)",
           all(r_.get("Cód. cliente") == ERP_CODE for r_ in rows),
           str([r_.get("Cód. cliente") for r_ in rows]))
        ck("EXCEL: no internal tenant-prefixed ids leak into any cell",
           not any(tenant + ":" in str(v) for r_ in rows for v in r_.values()), "")

    async with db.get_pool().acquire() as con:
        orow = await con.fetchrow(
            "select status, value_snapshot_cents, value_snapshot from orders where id=$1", oid)
    snap = (dict(orow).get("value_snapshot") if orow else None)
    snap = snap if isinstance(snap, dict) else _json.loads(snap or "{}")
    ck("SNAPSHOT: order marked sent + value frozen at dispatch (WP-4 meter input)",
       bool(orow) and orow["status"] == "sent" and orow["value_snapshot_cents"] is not None,
       f"cents={orow and orow['value_snapshot_cents']} unpriced={snap.get('unpriced_lines')}")
    if snap.get("unpriced_lines"):
        rough(f"{snap['unpriced_lines']} line(s) dispatched UNPRICED (snapshot flags them) — "
              f"chat orders rarely state prices; catalogue price fallback covered the rest")

    # the queue is clear now → the nudge tab must resolve
    c1 = await app_mod._review_nudge_scan(ST, now=now_real)
    async with db.control_pool().acquire() as con:
        t_status = await con.fetchval(
            "select status from open_loop where scope=$1::ltree and entity_key=$2", BOSS_SCOPE, ekey)
    ck("CLEAR: queue empty after dispatch → orders.review_clear resolves the tab",
       c1.get("cleared") == 1 and t_status == "resolved", f"c1={c1} tab={t_status}")

    # ═══ restore + cleanup ════════════════════════════════════════════════════════════════════════
    orders_mod.send_email_async = _real_email
    twilio_wa.send_whatsapp = _real_send
    if not KEEP:
        await _cleanup(tenant)
    else:
        print(f"  [KEEP] rows left in place (tenant={tenant}); fixtures in {tmpdir}")

    passed = sum(R)
    print(f"\nREHEARSAL DIRTY: {passed}/{len(R)} green — "
          f"{'THE UGLY TARIFA SURVIVED THE WHOLE LOOP' if passed == len(R) else 'STRUCTURAL FINDINGS ABOVE'}")
    if ROUGH:
        print("ROUGH EDGES OBSERVED:")
        for i, r_ in enumerate(ROUGH, 1):
            print(f"  {i}. {r_}")
    else:
        print("ROUGH EDGES OBSERVED: none")
    return 0 if passed == len(R) else 1


raise SystemExit(asyncio.run(main()))
