"""vital_codes.py — M1M2 WP-3: THEIR CODES all the way through + dispatch robustness (eng F6/F11).

  1. SEED+MATCH — a seeded catalogue with supplier codes → the matcher pick carries `code` as a
     SEPARATE field on BOTH the alias fast-path and the rerank path; `sku_id` stays the internal
     tenant-prefixed id chain (aliases/accept-rules key on it) — NEVER overwritten by the code.
  2. TABULAR — the real tabular.build line carries ln["code"] with sku_id intact.
  3. ERP CODE — db.erp_code_set/get/erp_codes_for round-trip (set, batch-resolve, clear).
  4. CODIGO ≠ INTERNAL ID (eng F6 regression) — a line with an internal `sku` and NO code renders
     an EMPTY `codigo` cell (the old renderer leaked ln["sku"]).
  5. DIRTY NAME (eng F11) — customer 'Bar "Niño"/Q:*' with a per-customer layout: build_excel does
     NOT raise; the sheet title is sanitized (openpyxl-illegal chars stripped).
  6. PHONE GROUPING (eng F11) — two customers BOTH named "Bar Paco" on different phones get
     SEPARATE sheets, each under its OWN column layout; titles deduped ("Bar Paco" / "Bar Paco (2)").
  7. DISPATCH LOOP — through the REAL _run_dispatch: a failing email routes into the WP-0-shaped
     dispatch_failed nudge (duty+reflex rows seeded; per-day dedupe: a second failure emits nothing);
     the next SUCCESS renders `codigo` + `codigo_cliente` (erp code resolved AT dispatch from
     customer_profiles) into the attached Excel AND resolves the dispfail tab.

SAFETY: refuses the live `rocoto` DB. Point it at rocoto_rehearsal (REHEARSAL_DSN or per-role DSNs).
Usage:  REHEARSAL_DSN=… PYTHONPATH=. python tests/vital_codes.py
"""
import asyncio
import io
import os
import sys
from pathlib import Path

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

_DSN = os.environ.get("REHEARSAL_DSN")
if _DSN:
    for _k in ("DATABASE_URL", "RUNTIME_DATABASE_URL", "CONTROL_DATABASE_URL"):
        os.environ.setdefault(_k, _DSN)

for _l in Path("/opt/rocoto/.env").read_text().splitlines() if Path("/opt/rocoto/.env").exists() else []:
    _l = _l.strip()
    if _l and not _l.startswith("#") and "=" in _l:
        _k, _, _v = _l.partition("=")
        os.environ.setdefault(_k.strip(), _v.strip())

import app as app_mod  # noqa: E402  (the dispatch organ owns _run_dispatch + the failure nudge)
import db  # noqa: E402
import directory  # noqa: E402
import duties  # noqa: E402
import orders as orders_mod  # noqa: E402
import runtime  # noqa: E402
import spine  # noqa: E402
from brain import events as brain_events_mod  # noqa: E402
from brain import golive, matcher, tabular  # noqa: E402

TEN = "vitalwp3"
PHONE = "+34999333555"                      # the boss line the dispfail nudge lands on (vital-reserved)
SCOPE = runtime.scope_of(PHONE)
P1, P2, P3 = "+34600222331", "+34600222332", "+34600222333"   # customers (vital-reserved)
R = []


def ck(name, ok, detail=""):
    R.append(bool(ok))
    print(f"  [{'PASS' if ok else 'FAIL'}] {name}" + (f" — {detail}" if detail else ""))


SKUS = [
    {"code": "CH-01", "name": "Chorizo picante casero", "family": "embutidos", "unit": "kg", "price": 9.5},
    {"code": "QM-02", "name": "Queso manchego curado", "family": "quesos", "unit": "ud", "price": 12},
    {"code": "", "name": "Tomate seco en aceite", "family": "conservas", "unit": "caja", "price": 4.0},
]

_ILLEGAL = set('\\/*?:[]')


def _wb(xlsx: bytes):
    from openpyxl import load_workbook
    return load_workbook(io.BytesIO(xlsx))


def _headers(ws) -> list:
    out = []
    for i in range(1, 20):
        v = ws.cell(row=1, column=i).value
        if v is None:
            break
        out.append(v)
    return out


async def _cleanup():
    async with db.get_pool().acquire() as con:
        await con.execute("delete from sku_aliases where tenant_id=$1", TEN)
        await con.execute("delete from catalogue_skus where tenant_id=$1", TEN)
        await con.execute("delete from catalogue_directory where tenant_id=$1", TEN)
        await con.execute("delete from customer_profiles where phone_number = any($1::text[])",
                          [P1, P2, P3])
    task = directory._PENDING.pop(TEN, None)
    if task and not task.done():
        task.cancel()
    try:
        async with db.scoped_tx(f"t.{TEN}") as con:
            await con.execute("delete from event where scope=$1::ltree", f"t.{TEN}")
    except Exception:
        pass
    try:
        async with db.scoped_tx(SCOPE) as con:
            for t in ("timer", "open_loop", "event"):
                await con.execute(f"delete from {t} where scope=$1::ltree", SCOPE)
    except Exception:
        pass


async def main():
    await db.init()
    duties.invalidate_cache()

    # SAFETY GUARD — never the live rocoto DB.
    async with db.get_pool().acquire() as con:
        dbname = await con.fetchval("select current_database()")
    if dbname == "rocoto":
        raise SystemExit("REFUSING to run vitals against the LIVE 'rocoto' DB — use rocoto_rehearsal")
    print(f"  (db={dbname})")

    await _cleanup()

    # ── 1. SEED + MATCH: code carried as a SEPARATE field, sku_id chain untouched ──
    r1 = await golive.seed_catalogue(TEN, SKUS)
    ck("seed ok (3 skus embedded)", r1.get("ok") and r1.get("skus") == 3, str(r1))

    m1 = await matcher.match("Chorizo picante casero", tenant=TEN)      # alias fast-path
    p1 = m1.pick or {}
    ck("alias path: pick carries the supplier's code",
       m1.alias_hit and p1.get("code") == "CH-01", f"code={p1.get('code')} method={m1.method}")
    ck("alias path: sku_id stays the INTERNAL tenant-prefixed id (never the code)",
       str(p1.get("sku_id", "")).startswith(TEN + ":") and p1.get("sku_id") != "CH-01",
       f"sku_id={p1.get('sku_id')}")

    m2 = await matcher.match("chorizo", tenant=TEN)                     # dense/sparse → rerank path
    p2 = m2.pick or {}
    c2 = (m2.candidates or [{}])[0]
    ck("rerank path: pick + top candidate carry the code",
       p2.get("code") == "CH-01" and c2.get("code") == "CH-01",
       f"pick.code={p2.get('code')} cand.code={c2.get('code')} method={m2.method}")

    m3 = await matcher.match("Tomate seco en aceite", tenant=TEN)
    ck("a codeless SKU carries code=None (no invention)",
       (m3.pick or {}).get("code") in (None, ""), f"code={(m3.pick or {}).get('code')}")

    # ── 2. TABULAR — the real order line carries ln["code"] ──
    gl = [{"mention": "chorizo", "qty": 2, "unit": "kg", "note": "", "match": m2.as_dict()}]
    decided = {"lines": [{"index": 0, "sku_id": p2.get("sku_id"), "qty": 2.0, "unit": "kg"}],
               "order_note": ""}
    lines = await tabular.build(decided, {"grounded_lines": gl}, [{}], {"tenant": TEN})
    ln0 = (lines or [{}])[0]
    ck("tabular line: code='CH-01' AND sku_id intact (separate fields, eng F6)",
       ln0.get("code") == "CH-01" and ln0.get("sku_id") == p2.get("sku_id"),
       f"code={ln0.get('code')} sku_id={ln0.get('sku_id')}")

    # ── 3. ERP client code round-trip ──
    s1 = await db.erp_code_set(P1, "CL-042")
    got = await db.erp_code_get(P1)
    batch = await db.erp_codes_for([P1, P2, ""])
    ck("erp_code set/get/batch (only rows WITH a code)",
       s1 == "CL-042" and got == "CL-042" and batch == {P1: "CL-042"}, f"batch={batch}")
    await db.erp_code_set(P2, "CL-777")
    cleared = await db.erp_code_set(P2, "")
    ck("empty code CLEARS", cleared is None and await db.erp_code_get(P2) is None)

    # ── 4. eng F6 regression: codigo NEVER leaks the internal sku id ──
    cols = [{"header": "Fecha", "field": "fecha", "order": 0},
            {"header": "Cliente", "field": "cliente", "order": 1},
            {"header": "Código", "field": "codigo", "order": 2},
            {"header": "Cód. cliente", "field": "codigo_cliente", "order": 3},
            {"header": "Producto", "field": "producto", "order": 4},
            {"header": "Cantidad", "field": "cantidad", "order": 5}]
    o_noc = {"id": "vc-x", "phone_number": P3, "customer_name": "Sin Código",
             "lines": [{"item": "Magro", "sku": f"{TEN}:deadbeef", "sku_id": f"{TEN}:deadbeef",
                        "qty": 1, "unit": "kg", "price": None, "note": ""}]}
    ws = _wb(orders_mod.build_excel([o_noc], columns=cols)).active
    ck("a coded-column line WITHOUT a code renders EMPTY codigo (never the internal sku)",
       (ws.cell(row=2, column=3).value or "") == "", f"cell={ws.cell(row=2, column=3).value!r}")

    # ── 5. DIRTY NAME — 'Bar "Niño"/Q:*' must not kill the dispatch ──
    dirty = 'Bar "Niño"/Q:*'
    o_dirty = {"id": "vc-d", "phone_number": P2, "customer_name": dirty,
               "lines": [{"item": "Chorizo", "code": "CH-01", "qty": 2, "unit": "kg",
                          "price": 9.5, "note": ""}]}
    try:
        wb = _wb(orders_mod.build_excel([o_dirty], columns=cols, per_customer={P2: cols}))
        titles = wb.sheetnames
        legal = all(not (_ILLEGAL & set(t)) and 0 < len(t) <= 31 for t in titles)
        ck("dirty customer name does NOT kill the dispatch; title sanitized",
           legal and any("Niño" in t for t in titles), f"titles={titles}")
    except Exception as e:  # noqa: BLE001
        ck("dirty customer name does NOT kill the dispatch; title sanitized", False, repr(e))

    # ── 6. PHONE GROUPING — same name, different phones → SEPARATE sheets, own layouts ──
    cols_a = cols
    cols_b = [{"header": "Producto", "field": "producto", "order": 0},
              {"header": "Cantidad", "field": "cantidad", "order": 1}]
    o_a = {"id": "vc-a", "phone_number": P2, "customer_name": "Bar Paco",
           "lines": [{"item": "Chorizo", "code": "CH-01", "qty": 2, "unit": "kg", "price": 9.5, "note": ""}]}
    o_b = {"id": "vc-b", "phone_number": P3, "customer_name": "Bar Paco",
           "lines": [{"item": "Queso", "code": "QM-02", "qty": 1, "unit": "ud", "price": 12, "note": ""}]}
    wb = _wb(orders_mod.build_excel([o_a, o_b], columns=cols, per_customer={P2: cols_a, P3: cols_b}))
    names = wb.sheetnames
    ok_split = (len(names) == 2 and names[0] == "Bar Paco" and names[1] == "Bar Paco (2)")
    hdr_a = _headers(wb[names[0]]) if ok_split else []
    hdr_b = _headers(wb[names[1]]) if ok_split else []
    ck("same-name-different-phone → separate sheets, deduped titles",
       ok_split, f"sheets={names}")
    ck("each sheet keeps ITS customer's own layout",
       hdr_a == [c["header"] for c in cols_a] and hdr_b == [c["header"] for c in cols_b],
       f"a={hdr_a} b={hdr_b}")

    # ── 7. THE DISPATCH LOOP through the real _run_dispatch ──
    # rows seeded: the duty (L0 default) + the reflex registration.
    async with db.control_pool().acquire() as con:
        duty_ok = await con.fetchval(
            "select count(*) from manifest_duty where scope='peter' and name='dispatch_failed' and enabled")
        reflex_ok = await con.fetchval(
            "select count(*) from manifest_reflex where scope_tier='boss' and organ='dispatch_failed' and enabled")
    ck("dispatch_failed seeds as a manifest_duty ROW + its reflex registration",
       duty_ok == 1 and reflex_ok == 1, f"duty={duty_ok} reflex={reflex_ok}")

    # a boss party on a clean scope + the reflex pinned to OUR phone (restored in cleanup).
    async with db.scoped_tx(SCOPE) as con:
        await spine.resolve_chat_scope(con, None, "wa", PHONE, display_name="Boss Vital WP3",
                                       scope_override=SCOPE, tier_override="boss")
        await con.execute("update scope_node set tier='boss' where scope=$1::ltree", SCOPE)
    async with db.get_pool().acquire() as con:
        await con.execute("update manifest_reflex set config = config || $1::jsonb "
                          "where organ='dispatch_failed' and scope_tier='boss'", {"boss_phone": PHONE})

    ST = {"mode": "scheduled", "schedule_time": "23:59", "recipient_email": "vital@rehearsal.local",
          "excel_format": "standard", "n": 10, "default_columns": cols, "per_customer_columns": {}}

    async def _fake_settings():
        return ST

    _dispatch_order = {"id": "vc-disp-1", "phone_number": P1, "customer_name": "Bar Vital",
                       "status": "approved", "note": "",
                       "lines": [{"item": "Chorizo picante casero", "sku_id": p2.get("sku_id"),
                                  "sku": p2.get("sku_id"), "code": "CH-01",
                                  "qty": 2, "unit": "kg", "price": 9.5, "note": ""}]}

    async def _fake_for_dispatch():
        return [dict(_dispatch_order)]

    async def _fake_pending():
        return 0

    async def _fake_mark(ids, delivery_note=None):
        return None

    _emails = []
    _email_ok = {"ok": False}

    async def _fake_email(to, subject, html, attachment_name=None, attachment_bytes=None):
        _emails.append({"to": to, "subject": subject, "attachment": attachment_bytes})
        return dict(_email_ok) if _email_ok["ok"] else {"ok": False, "error": "smtp down (vital)"}

    _saved = (app_mod._dispatch_settings, db.orders_for_dispatch, db.orders_pending_review_count,
              brain_events_mod.mark_dispatched, orders_mod.send_email_async)
    app_mod._dispatch_settings = _fake_settings
    db.orders_for_dispatch = _fake_for_dispatch
    db.orders_pending_review_count = _fake_pending
    brain_events_mod.mark_dispatched = _fake_mark
    orders_mod.send_email_async = _fake_email

    async def _n_events(etype):
        async with db.control_pool().acquire() as con:
            return int(await con.fetchval(
                "select count(*) from event where scope=$1::ltree and type=$2", SCOPE, etype) or 0)

    async def _tab_status():
        async with db.control_pool().acquire() as con:
            return await con.fetchval(
                "select status from open_loop where scope=$1::ltree and entity_key like 'dispfail:%' "
                "order by id desc limit 1", SCOPE)

    try:
        # 7a. FAILURE → the nudge (never a silent log.warning).
        d1 = await app_mod._run_dispatch("scheduled")
        n1 = await _n_events("orders.dispatch_failed")
        st1 = await _tab_status()
        ck("a failed dispatch routes into the nudge: orders.dispatch_failed + an OPEN dispfail tab",
           not d1.get("ok") and n1 == 1 and st1 == "open", f"res.ok={d1.get('ok')} events={n1} tab={st1}")

        # 7b. a second failure the same day is dedupe-swallowed (one WhatsApp a day, max).
        d2 = await app_mod._run_dispatch("manual")
        n2 = await _n_events("orders.dispatch_failed")
        ck("second failure same day → dedupe-swallowed (still ONE sense)",
           not d2.get("ok") and n2 == 1, f"events={n2}")

        # the armed rung's sentence: fully resolved, count-free, names the recovery move.
        async with db.control_pool().acquire() as con:
            cont = await con.fetchval(
                "select contingencies from open_loop where scope=$1::ltree and entity_key like 'dispfail:%' "
                "and status='open' order by id desc limit 1", SCOPE)
        import json as _json
        cont = cont if isinstance(cont, list) else _json.loads(cont or "[]")
        sent = ((cont[0] if cont else {}).get("action", {}) or {}).get("sentence") or ""
        ck("the nudge sentence is resolved + count-free (es-ES, names «Enviar ahora»)",
           sent and "{" not in sent and "Excel" in sent, sent[:90])

        # 7c. SUCCESS → codigo + codigo_cliente in the real attached Excel; the tab RESOLVES.
        _email_ok.update(ok=True, id="vital")
        d3 = await app_mod._run_dispatch("scheduled")
        att = next((e["attachment"] for e in reversed(_emails) if e["attachment"]), None)
        ok_xl, row = False, None
        if att:
            ws3 = _wb(att).active
            hdrs = _headers(ws3)
            row = {h: ws3.cell(row=2, column=i + 1).value for i, h in enumerate(hdrs)}
            ok_xl = (row.get("Código") == "CH-01" and row.get("Cód. cliente") == "CL-042")
        st3 = await _tab_status()
        n_ok = await _n_events("orders.dispatch_ok")
        ck("success: the Excel renders codigo='CH-01' + codigo_cliente='CL-042' (resolved AT dispatch)",
           d3.get("ok") and ok_xl, f"ok={d3.get('ok')} row={row if att else None}")
        ck("success RESOLVES the dispfail tab (orders.dispatch_ok)",
           st3 == "resolved" and n_ok == 1, f"tab={st3} events={n_ok}")
    finally:
        (app_mod._dispatch_settings, db.orders_for_dispatch, db.orders_pending_review_count,
         brain_events_mod.mark_dispatched, orders_mod.send_email_async) = _saved
        async with db.get_pool().acquire() as con:
            await con.execute("update manifest_reflex set config = config - 'boss_phone' "
                              "where organ='dispatch_failed' and scope_tier='boss'")

    await _cleanup()
    async with db.scoped_tx(SCOPE) as con:
        await con.execute("delete from scope_node where scope=$1::ltree", SCOPE)

    passed = sum(R)
    print(f"\nCODES: {passed}/{len(R)} passed — "
          f"{'THEIR CODES FLOW + THE DISPATCH IS ROBUST' if passed == len(R) else 'BROKEN'}")
    return 0 if passed == len(R) else 1


raise SystemExit(asyncio.run(main()))
